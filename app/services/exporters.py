"""Turn report rows into CSV, Excel (.xlsx) and PDF files.

Security: text starting with = + - @ (or a tab/carriage return) is prefixed with an apostrophe
in CSV and Excel output, so a malicious value such as a vehicle named "=HYPERLINK(...)" can't
run as a formula when the file is opened (OWASP "CSV injection").
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from flask import current_app

from app.services.reports import Column, ReportData

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


@dataclass
class ReportMeta:
    title: str
    period: str
    generated_by: str
    generated_at: datetime  # naive UTC


def _local(value: datetime) -> datetime:
    tz = ZoneInfo(current_app.config["APP_TIMEZONE"])
    return value.replace(tzinfo=UTC).astimezone(tz).replace(tzinfo=None)


def safe_text(value: str) -> str:
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def format_value(column: Column, value) -> str:
    """Human-readable cell text (preview and PDF)."""
    if value is None or value == "":
        return "—" if column.kind != "text" else ""
    match column.kind:
        case "int":
            return f"{value:,.0f}"
        case "float1":
            return f"{value:,.1f}"
        case "float2":
            return f"{value:,.2f}"
        case "inr":
            return f"₹{value:,.0f}"
        case "inr2":
            return f"₹{value:,.2f}"
        case "pct":
            return f"{value:+.1f}%"
        case "date":
            return value.strftime("%d %b %Y")
        case "datetime":
            return _local(value).strftime("%d %b %Y %H:%M")
        case _:
            return str(value)


def _raw(column: Column, value):
    """Machine-friendly value (CSV and Excel): numbers stay numbers, times in local time."""
    if value is None:
        return None
    if column.kind == "datetime":
        return _local(value)
    if isinstance(value, str):
        return safe_text(value)
    return value


# --- CSV --------------------------------------------------------------------------------------


def to_csv(columns: list[Column], data: ReportData) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([c.label for c in columns])
    for row in data.rows:
        out = []
        for c in columns:
            value = _raw(c, row.get(c.key))
            if isinstance(value, datetime):
                value = value.strftime("%Y-%m-%d %H:%M")
            elif isinstance(value, float):
                value = round(value, 4)
            out.append("" if value is None else value)
        writer.writerow(out)
    # UTF-8 with BOM so Excel shows ₹ and other non-ASCII text correctly.
    return buffer.getvalue().encode("utf-8-sig")


# --- Excel ------------------------------------------------------------------------------------

_XLSX_FORMATS = {
    "int": "#,##0",
    "float1": "#,##0.0",
    "float2": "#,##0.00",
    "inr": '"₹"#,##0',
    "inr2": '"₹"#,##0.00',
    "pct": '+0.0"%";-0.0"%";0.0"%"',
    "date": "dd mmm yyyy",
    "datetime": "dd mmm yyyy hh:mm",
}


def to_xlsx(columns: list[Column], data: ReportData, meta: ReportMeta) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    header_fill = PatternFill("solid", fgColor="2E7D32")
    ws.append([c.label for c in columns])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row in data.rows:
        ws.append([_raw(c, row.get(c.key)) for c in columns])
    for idx, column in enumerate(columns, start=1):
        letter = get_column_letter(idx)
        fmt = _XLSX_FORMATS.get(column.kind)
        if fmt:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = fmt
        longest = max(
            [len(column.label)]
            + [len(format_value(column, r.get(column.key))) for r in data.rows[:500]]
        )
        ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)
    ws.freeze_panes = "A2"
    if data.rows:
        ws.auto_filter.ref = ws.dimensions

    info = wb.create_sheet("Summary")
    info.append([meta.title])
    info["A1"].font = Font(bold=True, size=14)
    info.append(["Period", meta.period])
    info.append(["Generated", f"{_local(meta.generated_at):%d %b %Y %H:%M} by {meta.generated_by}"])
    info.append(["Rows", len(data.rows)])
    info.append([])
    for label, value in data.summary:
        info.append([label, value])
    if data.note:
        info.append([])
        info.append(["Note", data.note])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- PDF --------------------------------------------------------------------------------------


def to_pdf(columns: list[Column], data: ReportData, meta: ReportMeta) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table

    font, bold = _register_unicode_font(pdfmetrics, TTFont)
    pagesize = landscape(A4) if len(columns) > 6 else A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=meta.title,
        author="EV Fleet Monitor",
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle(
        "h1", parent=styles["Heading1"], fontName=bold, textColor=colors.HexColor("#2E7D32")
    )
    small = ParagraphStyle("small", parent=styles["Normal"], fontName=font, fontSize=8, leading=10)
    cell = ParagraphStyle("cell", parent=small, fontSize=7, leading=8.5)
    head = ParagraphStyle("head", parent=cell, fontName=bold, textColor=colors.white)

    story = [
        Paragraph(meta.title, h1),
        Paragraph(
            f"Period: {meta.period} · Generated {_local(meta.generated_at):%d %b %Y %H:%M} "
            f"by {meta.generated_by} · {len(data.rows)} row{'' if len(data.rows) == 1 else 's'}",
            small,
        ),
        Spacer(1, 4 * mm),
    ]
    if data.summary:
        summary = Table(
            [[label for label, _ in data.summary], [value for _, value in data.summary]],
            hAlign="LEFT",
        )
        summary.setStyle(
            [
                ("FONTNAME", (0, 0), (-1, 0), font),
                ("FONTNAME", (0, 1), (-1, 1), bold),
                ("FONTSIZE", (0, 0), (-1, 0), 7),
                ("FONTSIZE", (0, 1), (-1, 1), 10),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.grey),
                ("RIGHTPADDING", (0, 0), (-1, -1), 14),
            ]
        )
        story += [summary, Spacer(1, 5 * mm)]
    if data.note:
        story += [Paragraph(data.note, small), Spacer(1, 3 * mm)]

    if data.rows:
        header = [Paragraph(_escape(c.label), head) for c in columns]
        body = [
            [Paragraph(_escape(format_value(c, r.get(c.key))), cell) for c in columns]
            for r in data.rows
        ]
        table = LongTable([header] + body, repeatRows=1, hAlign="LEFT")
        table.setStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E7D32")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F8E9")]),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C8E6C9")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
        story.append(table)
    else:
        story.append(Paragraph("No data for this period and selection.", small))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont(font, 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 8 * mm, f"EV Fleet Monitor · {meta.title}")
        canvas.drawRightString(pagesize[0] - 12 * mm, 8 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _escape(text: str) -> str:
    """ReportLab paragraphs use XML-like markup, so escape special characters."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _register_unicode_font(pdfmetrics, TTFont) -> tuple[str, str]:  # noqa: N803
    """Register the bundled DejaVu Sans font (returns regular and bold names).

    The PDF built-in fonts (Helvetica etc.) have no ₹ glyph. DejaVu Sans does, is free to
    redistribute (licence in app/static/fonts/), and bundling it makes PDFs identical on
    Windows, Docker and the hosting server.
    """
    fonts = Path(current_app.static_folder) / "fonts"
    if "DejaVuSans" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVuSans", fonts / "DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", fonts / "DejaVuSans-Bold.ttf"))
    return "DejaVuSans", "DejaVuSans-Bold"


def file_stamp(first: date, last: date) -> str:
    return f"{first:%Y%m%d}-{last:%Y%m%d}"
