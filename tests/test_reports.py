import csv
import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.extensions import db
from app.models import (
    AlertSeverity,
    AlertType,
    DailyEnergy,
    ServiceRecord,
    ServiceType,
    Telemetry,
)
from app.services.alerts import raise_alert
from app.services.energy import local_today
from app.services.exporters import safe_text
from app.utils import utcnow


@pytest.fixture
def fleet(app, make_vehicle, make_user):
    """Two vehicles with energy, service, alert and driving data for yesterday."""
    driver = make_user(email="ravi@example.com", name="Ravi Kumar")
    v1, _ = make_vehicle(plate="DL01AB0001", driver=driver)
    # A malicious make: the "Vehicle" cell then starts with "=" (a spreadsheet formula).
    v2, _ = make_vehicle(plate="DL01AB0002", make='=HYPERLINK("http://evil")', model="X")
    day = local_today() - timedelta(days=1)
    for v, km in ((v1, 120.0), (v2, 80.0)):
        db.session.add(
            DailyEnergy(
                vehicle_id=v.id,
                day=day,
                distance_km=km,
                energy_used_kwh=km / 7,
                grid_energy_kwh=km / 6,
                energy_cost_inr=km * 1.6,
                charging_cost_inr=km * 1.5,
                tariff_inr_per_kwh=10,
            )
        )
    db.session.add(
        ServiceRecord(
            vehicle_id=v1.id,
            service_type=ServiceType.ROUTINE,
            service_date=day,
            cost_inr=2500,
            description="Annual service",
        )
    )
    raise_alert(v2, AlertType.LOW_BATTERY, AlertSeverity.CRITICAL, "Battery at 8%")
    now = utcnow()
    for i, (speed, accel, odo) in enumerate([(40, 0.1, 0), (40, -5.0, 30), (40, 0.1, 60)]):
        db.session.add(
            Telemetry(
                vehicle_id=v1.id,
                driver_id=driver.id,
                recorded_at=now - timedelta(minutes=30 - i),
                lat=1,
                lon=1,
                speed_kmh=speed,
                soc_pct=80 - i,
                acceleration_mps2=accel,
                odometer_km=odo,
            )
        )
    db.session.commit()
    return {"driver": driver, "v1": v1, "v2": v2, "day": day}


def _get(client, path="/reports/", **params):
    return client.get(path, query_string=params)


def test_builder_page_without_params(client, manager):
    resp = client.get("/reports/")
    assert resp.status_code == 200 and b"Choose a report and press Preview" in resp.data


@pytest.mark.parametrize(
    ("report", "expected"),
    [
        ("fleet_summary", [b"DL01AB0001", b"DL01AB0002", "₹2,692".encode()]),
        ("energy_daily", [b"DL01AB0001", b"120.0"]),
        ("driver_behaviour", [b"Ravi Kumar", b"Good"]),  # 1 harsh brake in 60 km -> 93
        ("alerts", [b"Low battery", b"Battery at 8%"]),
        ("service_history", [b"Annual service", "₹2,500".encode()]),
    ],
)
def test_preview_each_report(client, manager, fleet, report, expected):
    resp = _get(client, type=report)
    assert resp.status_code == 200
    for needle in expected:
        assert needle in resp.data, needle


def test_column_and_vehicle_selection(client, manager, fleet):
    resp = _get(client, type="fleet_summary", col=["plate", "distance_km"], vehicle=fleet["v1"].id)
    table = resp.data.split(b"<tbody>")[1].split(b"</tbody>")[0]  # the report, not the form
    assert b"DL01AB0001" in table and b"DL01AB0002" not in table
    header = resp.data.split(b"<thead")[1].split(b"</thead>")[0]
    assert b"Distance (km)" in header and b"Energy cost" not in header


def test_csv_download(client, manager, fleet):
    resp = _get(client, "/reports/download/csv", type="energy_daily")
    assert resp.status_code == 200 and resp.content_type == "text/csv; charset=utf-8"
    assert "attachment;" in resp.headers["Content-Disposition"]
    assert resp.data.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel
    rows = list(csv.reader(io.StringIO(resp.data.decode("utf-8-sig"))))
    assert rows[0][:3] == ["Date", "Registration", "Distance (km)"]
    assert len(rows) == 3


def test_xlsx_download_is_formatted(client, manager, fleet):
    resp = _get(client, "/reports/download/xlsx", type="fleet_summary", title="Board pack")
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb["Data"]
    assert ws["A1"].value == "Registration" and ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    distance_col = [c.value for c in ws[1]].index("Distance (km)") + 1
    assert isinstance(ws.cell(row=2, column=distance_col).value, int | float)  # not text
    summary = wb["Summary"]
    assert summary["A1"].value == "Board pack"
    assert any(row[0].value == "Operating cost" for row in summary.iter_rows())


def test_pdf_download(client, manager, fleet):
    for report in ("fleet_summary", "alerts", "service_history"):
        resp = _get(client, "/reports/download/pdf", type=report)
        assert resp.status_code == 200 and resp.mimetype == "application/pdf"
        assert resp.data.startswith(b"%PDF") and len(resp.data) > 2000


def test_empty_report_exports_cleanly(client, manager):
    for fmt in ("csv", "xlsx", "pdf"):
        assert _get(client, f"/reports/download/{fmt}", type="alerts").status_code == 200


def test_spreadsheet_formula_injection_is_neutralised(client, manager, fleet):
    resp = _get(client, "/reports/download/csv", type="fleet_summary")
    assert "'=HYPERLINK" in resp.data.decode("utf-8-sig")
    wb = load_workbook(
        io.BytesIO(_get(client, "/reports/download/xlsx", type="fleet_summary").data)
    )
    values = [c.value for row in wb["Data"].iter_rows() for c in row]
    assert '\'=HYPERLINK("http://evil") X' in values
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)


@pytest.mark.parametrize(
    ("raw", "safe"),
    [("=1+1", "'=1+1"), ("+91", "'+91"), ("-5", "'-5"), ("@SUM", "'@SUM"), ("Tata", "Tata")],
)
def test_safe_text(raw, safe):
    assert safe_text(raw) == safe


def test_driver_reports_are_limited_to_own_data(client, fleet, login):
    login(email="ravi@example.com")
    resp = _get(client, type="fleet_summary", vehicle=fleet["v2"].id)  # not theirs: ignored
    assert b"DL01AB0001" in resp.data and b"DL01AB0002" not in resp.data
    resp = _get(client, type="alerts")
    assert b"Battery at 8%" not in resp.data  # alert is on someone else's vehicle
    resp = _get(client, type="driver_behaviour")
    assert b"Ravi Kumar" in resp.data


def test_date_validation(client, manager):
    today = local_today()
    resp = _get(client, type="alerts", start="2026-13-40")
    assert b"YYYY-MM-DD" in resp.data
    resp = _get(
        client, type="alerts", start=today.isoformat(), end=(today - timedelta(days=5)).isoformat()
    )
    assert b"start date must be on or before" in resp.data
    resp = _get(client, type="alerts", start=(today - timedelta(days=800)).isoformat())
    assert b"at most 366 days" in resp.data
    future = (today + timedelta(days=10)).isoformat()
    resp = _get(client, type="alerts", end=future)
    assert today.strftime("%d %b %Y").encode() in resp.data  # clamped to today


def test_unknown_format_and_login_required(client, manager):
    assert client.get("/reports/download/docx").status_code == 404
    client.post("/logout")
    assert client.get("/reports/").status_code == 302


def test_row_limit(client, manager, fleet, monkeypatch):
    monkeypatch.setattr("app.reports.routes.MAX_ROWS", 1)
    resp = _get(client, "/reports/download/csv", type="energy_daily")
    assert resp.status_code == 413


def test_filename_reflects_report_and_period(client, manager):
    day = date(2026, 9, 1)
    resp = _get(
        client,
        "/reports/download/csv",
        type="service_history",
        start=day.isoformat(),
        end="2026-09-30",
    )
    assert (
        'filename="ev-fleet-service-history-20260901-20260930.csv"'
        in resp.headers["Content-Disposition"]
    )
